"""Bounded, source-addressable parsing; parser exceptions are terminal failures."""

import csv
import json
import re
import zipfile
from dataclasses import dataclass
from io import BytesIO, StringIO
from pathlib import PurePath


class InvalidFile(ValueError):
    pass


@dataclass(frozen=True)
class Block:
    text: str
    location: dict
    warning: str = ""


def _zip_guard(data: bytes, expected: str):
    if not data.startswith(b"PK"):
        raise InvalidFile("Invalid Office file signature")
    with zipfile.ZipFile(BytesIO(data)) as archive:
        if expected not in archive.namelist():
            raise InvalidFile("File format does not match extension")
        if sum(x.file_size for x in archive.infolist()) > 200 * 1024 * 1024:
            raise InvalidFile("Expanded archive exceeds 200 MiB")
        if any(x.flag_bits & 1 for x in archive.infolist()):
            raise InvalidFile("Encrypted archive is not supported")


def _ocr(image):
    import pytesseract

    if image.width * image.height > 40_000_000:
        raise InvalidFile("Image dimensions exceed OCR limit")
    return pytesseract.image_to_string(image, lang="chi_sim+eng", timeout=60)


def parse(data: bytes, filename: str) -> list[Block]:
    if not data or len(data) > 50 * 1024 * 1024:
        raise InvalidFile("File is empty or exceeds 50 MiB")
    ext = PurePath(filename).suffix.lower()
    blocks = []
    if ext in {".txt", ".md", ".csv"}:
        try:
            text = data.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise InvalidFile("Text must be UTF-8") from exc
        if "\x00" in text:
            raise InvalidFile("Binary input disguised as text")
        if ext == ".csv":
            rows = csv.reader(StringIO(text))
            header = next(rows, [])
            for row_num, row in enumerate(rows, 2):
                if row_num > 100_001:
                    raise InvalidFile("Table exceeds 100000 rows")
                blocks.append(
                    Block(
                        " | ".join(
                            f"{header[i] if i < len(header) else i + 1}: {value}"
                            for i, value in enumerate(row)
                        ),
                        {
                            "type": "table",
                            "sheet": "CSV",
                            "row": row_num,
                            "columns": len(row),
                        },
                    )
                )
        elif ext == ".md":
            headings = list(re.finditer(r"^#{1,6}[^\S\n]+[^\n]+", text, re.MULTILINE))
            starts = sorted(set([0] + [m.start() for m in headings]))
            path = []
            for index, start in enumerate(starts):
                end = starts[index + 1] if index + 1 < len(starts) else len(text)
                raw = text[start:end]
                heading = re.match(r"(#{1,6})[^\S\n]+([^\n]+)", raw)
                if heading:
                    depth = len(heading.group(1))
                    path = path[: depth - 1] + [heading.group(2)]
                blocks.append(
                    Block(
                        raw,
                        {
                            "type": "text",
                            "start": start,
                            "end": end,
                            "title_path": list(path),
                        },
                    )
                )
        else:
            for match in re.finditer(r"[^\n]+(?:\n(?!\n)[^\n]+)*", text):
                blocks.append(
                    Block(
                        match.group(),
                        {"type": "text", "start": match.start(), "end": match.end()},
                    )
                )
    elif ext == ".docx":
        _zip_guard(data, "word/document.xml")
        from docx import Document

        doc = Document(BytesIO(data))
        for i, paragraph in enumerate(doc.paragraphs):
            if paragraph.text.strip():
                blocks.append(
                    Block(paragraph.text, {"type": "paragraph", "paragraph": i + 1})
                )
        for index, table in enumerate(doc.tables):
            header = [cell.text for cell in table.rows[0].cells]
            for row_num, row in enumerate(table.rows[1:], 2):
                blocks.append(
                    Block(
                        " | ".join(
                            f"{header[i]}: {cell.text}"
                            for i, cell in enumerate(row.cells)
                        ),
                        {"type": "table", "table": index + 1, "row": row_num},
                    )
                )
    elif ext == ".xlsx":
        _zip_guard(data, "xl/workbook.xml")
        from openpyxl import load_workbook

        wb = load_workbook(
            BytesIO(data), read_only=True, data_only=True, keep_links=False
        )
        try:
            for sheet in wb:
                rows = sheet.iter_rows(values_only=True)
                header = next(rows, ())
                for row_num, row in enumerate(rows, 2):
                    if row_num > 100_001:
                        raise InvalidFile("Table exceeds 100000 rows")
                    blocks.append(
                        Block(
                            " | ".join(
                                f"{header[i] if i < len(header) else i + 1}: {value}"
                                for i, value in enumerate(row)
                                if value is not None
                            ),
                            {"type": "table", "sheet": sheet.title, "row": row_num},
                        )
                    )
        finally:
            wb.close()
    elif ext == ".pdf":
        if not data.startswith(b"%PDF-"):
            raise InvalidFile("Invalid PDF signature")
        from pypdf import PdfReader

        pdf = PdfReader(BytesIO(data))
        if pdf.is_encrypted or len(pdf.pages) > 500:
            raise InvalidFile("Encrypted PDF or PDF exceeds 500 pages")
        rendered = None
        try:
            for i, page in enumerate(pdf.pages):
                text = page.extract_text() or ""
                warning = ""
                if len(text.strip()) < 30:
                    import pypdfium2

                    if rendered is None:
                        rendered = pypdfium2.PdfDocument(data)
                    rendered_page = rendered[i]
                    bitmap = rendered_page.render(scale=1.5)
                    image = bitmap.to_pil()
                    try:
                        text = _ocr(image)
                    finally:
                        image.close()
                        bitmap.close()
                        rendered_page.close()
                    warning = "OCR 结果需人工核对"
                blocks.append(Block(text, {"type": "pdf", "page": i + 1}, warning))
        finally:
            if rendered is not None:
                rendered.close()
    elif ext in {".png", ".jpg", ".jpeg"}:
        from PIL import Image

        with Image.open(BytesIO(data)) as image:
            expected = "PNG" if ext == ".png" else "JPEG"
            if image.format != expected:
                raise InvalidFile("Image format does not match extension")
            blocks.append(
                Block(_ocr(image), {"type": "image", "page": 1}, "OCR 结果需人工核对")
            )
    else:
        raise InvalidFile("Unsupported file type")
    blocks = [b for b in blocks if b.text.strip()]
    if not blocks:
        raise InvalidFile("No readable content; inspect OCR or document structure")
    if sum(len(b.text) for b in blocks) > 10_000_000:
        raise InvalidFile("Extracted text exceeds processing limit")
    return blocks


def chunk(blocks: list[Block], target=400, maximum=600, overlap=60) -> list[dict]:
    import tiktoken

    if not 0 <= overlap < target <= maximum <= 600:
        raise ValueError("Invalid token configuration")
    encoding = tiktoken.get_encoding("cl100k_base")
    output = []
    for block in blocks:
        # Split on Unicode character boundaries; encoding budgets count actual token IDs.
        start = 0
        while start < len(block.text):
            lo, hi = start + 1, len(block.text)
            while lo < hi:
                mid = (lo + hi + 1) // 2
                if len(encoding.encode(block.text[start:mid])) <= target:
                    lo = mid
                else:
                    hi = mid - 1
            end = lo
            if end < len(block.text):
                cut = max(
                    block.text.rfind("\n", start, end),
                    block.text.rfind("。", start, end),
                    block.text.rfind(". ", start, end),
                )
                if cut > start + (end - start) // 2:
                    end = cut + 1
            raw = block.text[start:end]
            cleaned = re.sub(r"[ \t]+", " ", raw).strip()
            if cleaned:
                location = {
                    **block.location,
                    "block_start": start,
                    "block_end": end,
                    "warning": block.warning
                    or (
                        "切片较短，请检查上下文是否完整"
                        if len(encoding.encode(cleaned)) < 20
                        else ""
                    ),
                }
                output.append(
                    {
                        "source_text": raw,
                        "content": cleaned,
                        "token_count": len(encoding.encode(cleaned)),
                        "location": json.dumps(location, ensure_ascii=False),
                    }
                )
            if end >= len(block.text):
                break
            next_start = end
            while (
                next_start > start + 1
                and len(encoding.encode(block.text[next_start - 1 : end])) <= overlap
            ):
                next_start -= 1
            start = max(start + 1, next_start)
    return output
