import json
import uuid
from io import BytesIO

import pytest

from careflow import models
from careflow.parsing import Block, InvalidFile, chunk, parse
from careflow.retrieval import collection, rrf


def test_text_offsets_are_real():
    text = "# 操作指南\n设备型号 CF-123。\n\n错误码 E404：检查连接。"
    blocks = parse(text.encode(), "guide.md")
    assert all(text[b.location["start"] : b.location["end"]] == b.text for b in blocks)
    assert "CF-123" in "".join(c["content"] for c in chunk(blocks))


def test_unicode_chunks_preserve_characters_and_token_budget():
    import tiktoken

    text = "知识库检索😀型号CF-123错误码E404。" * 150
    chunks = chunk([Block(text, {"type": "text"})])
    enc = tiktoken.get_encoding("cl100k_base")
    assert len(chunks) > 2
    for c in chunks:
        loc = json.loads(c["location"])
        assert c["source_text"] == text[loc["block_start"] : loc["block_end"]]
        assert "�" not in c["content"]
        assert len(enc.encode(c["content"])) <= 600


def test_csv_retains_headers_and_rows():
    blocks = parse("型号,错误码,处理\nCF-1,E404,重新连接\n".encode(), "codes.csv")
    assert blocks[0].location["row"] == 2
    assert blocks[0].text == "型号: CF-1 | 错误码: E404 | 处理: 重新连接"


def test_docx_parser_extracts_paragraphs_and_tables():
    from docx import Document

    doc = Document()
    doc.add_paragraph("产品手册")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "型号"
    table.cell(0, 1).text = "说明"
    table.cell(1, 0).text = "CF-1"
    table.cell(1, 1).text = "测试设备"
    stream = BytesIO()
    doc.save(stream)
    blocks = parse(stream.getvalue(), "guide.docx")
    assert blocks[0].location["paragraph"] == 1
    assert "型号: CF-1" in blocks[1].text


def test_xlsx_preserves_sheet_coordinates():
    from openpyxl import Workbook

    wb = Workbook()
    sheet = wb.active
    sheet.title = "故障码"
    sheet.append(["编号", "说明"])
    sheet.append(["E404", "连接失败"])
    stream = BytesIO()
    wb.save(stream)
    blocks = parse(stream.getvalue(), "codes.xlsx")
    assert blocks[0].location["sheet"] == "故障码"
    assert blocks[0].location["cell_range"] == "A2:B2"
    assert "E404" in blocks[0].text


@pytest.mark.parametrize(
    "data,name",
    [
        (b"fake", "x.pdf"),
        (b"fake", "x.docx"),
        (b"\x00exe", "x.txt"),
        (b"data", "x.exe"),
        (b"", "empty.txt"),
    ],
)
def test_rejects_invalid_files(data, name):
    with pytest.raises(InvalidFile):
        parse(data, name)


def test_encrypted_pdf_rejected():
    from pypdf import PdfWriter

    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    writer.encrypt("secret")
    stream = BytesIO()
    writer.write(stream)
    with pytest.raises(InvalidFile):
        parse(stream.getvalue(), "encrypted.pdf")


def test_rrf_is_rank_fusion_not_model_rerank():
    result = rrf([{"id": "a"}, {"id": "b"}], [{"id": "b"}, {"id": "c"}])
    assert result[0]["id"] == "b"
    assert result[0]["score"] == pytest.approx(1 / 61 + 1 / 62)


def test_model_identity_includes_revision_not_just_dimensions(monkeypatch):
    monkeypatch.setenv("EMBEDDING_MODEL", "real-model")
    monkeypatch.setenv("EMBEDDING_REVISION", "revision-1")
    monkeypatch.setenv("EMBEDDING_DIMENSIONS", "1024")
    tenant = str(uuid.uuid4())
    first = collection(tenant)
    monkeypatch.setenv("EMBEDDING_REVISION", "revision-2")
    assert collection(tenant) != first


def test_missing_models_fail_without_fake_fallback(monkeypatch):
    monkeypatch.delenv("EMBEDDING_MODEL", raising=False)
    monkeypatch.delenv("EMBEDDING_BASE_URL", raising=False)
    with pytest.raises(models.ModelUnavailable):
        models.embed(["hello"])


def test_worker_requires_internal_identity(monkeypatch):
    from fastapi.testclient import TestClient

    from careflow.api import app

    monkeypatch.setenv("INTERNAL_TOKEN", "a" * 32)
    client = TestClient(app)
    result = client.post(
        "/internal/v1/recall",
        json={"tenant_id": str(uuid.uuid4()), "version_ids": [], "query": "q"},
    )
    assert result.status_code == 401


def test_empty_scope_never_queries_milvus(monkeypatch):
    import careflow.retrieval as retrieval

    def forbidden():
        raise AssertionError("Empty authorization scope must never search")

    monkeypatch.setattr(retrieval, "client", forbidden)
    assert retrieval.recall(str(uuid.uuid4()), [], "q")["fused"] == []


def test_evaluation_metrics_use_unique_retrieved_ids():
    from careflow.evaluation import metrics

    assert metrics(["a", "b"], ["c", "b", "b"], k=3) == {
        "recall": 0.5,
        "hit": 1,
        "mrr": 0.5,
        "empty_correct": None,
    }
    assert metrics([], [], k=6)["empty_correct"] is True


def test_markdown_sections_keep_headings_with_body():
    text = "# 手册\n总览\n\n## 连接失败\n先检查网络。\n\n## 重试\n请等待十秒。"
    blocks = parse(text.encode(), "guide.md")
    assert len(blocks) == 3
    assert blocks[1].location["title_path"] == ["手册", "连接失败"]
    assert "## 连接失败\n先检查网络。" in blocks[1].text
    assert all(text[b.location["start"] : b.location["end"]] == b.text for b in blocks)
