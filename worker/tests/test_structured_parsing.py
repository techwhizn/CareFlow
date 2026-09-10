"""Synthetic format fixtures exercise real parser libraries, no OCR/model substitutes."""

import json
from io import BytesIO

from careflow.parsing import Block, chunk, parse
from careflow.text_parsing import repeated_page_margins


def test_word_keeps_body_order_heading_path_lists_and_table_coordinates():
    from docx import Document

    document = Document()
    document.add_heading("Manual", level=1)
    document.add_paragraph("First instruction", style="List Bullet")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Model"
    table.cell(0, 1).text = "Weight (kg)"
    table.cell(1, 0).text = "CF-100"
    table.cell(1, 1).text = "12"
    document.add_paragraph("After table")
    stream = BytesIO()
    document.save(stream)
    blocks = parse(stream.getvalue(), "manual.docx")
    assert [b.location["type"] for b in blocks] == [
        "paragraph",
        "paragraph",
        "table",
        "paragraph",
    ]
    assert blocks[1].location["kind"] == "list"
    assert blocks[2].location["title_path"] == ["Manual"]
    assert blocks[2].location["column_end"] == 2
    assert "Weight (kg): 12" in blocks[2].text
    assert blocks[3].location["paragraph"] == 3


def test_merged_word_table_reports_warning_in_chunk_location():
    from docx import Document

    document = Document()
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).merge(table.cell(0, 1)).text = "Merged"
    table.cell(1, 0).text = "A"
    table.cell(1, 1).text = "B"
    stream = BytesIO()
    document.save(stream)
    blocks = parse(stream.getvalue(), "merged.docx")
    assert "合并" in blocks[0].warning
    assert "合并" in json.loads(chunk(blocks)[0]["location"])["warning"]


def test_fenced_markdown_heading_is_content_and_skipped_levels_do_not_become_siblings():
    text = "# Root\nIntro\n```text\n# fake heading\n```\n### First\nBody\n### Second\nNext\n"
    blocks = parse(text.encode(), "guide.md")
    assert len(blocks) == 3
    assert "# fake heading" in blocks[0].text
    assert blocks[2].location["title_path"] == ["Root", "Second"]
    for block in blocks:
        assert text[block.location["start"] : block.location["end"]] == block.text


def test_excel_formulas_are_visible_uncomputed_and_coordinates_include_empty_columns():
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Measurements"
    sheet.append(["Model", "Amount", "Total"])
    sheet.append(["CF-100", None, "=1+2"])
    stream = BytesIO()
    workbook.save(stream)
    block = parse(stream.getvalue(), "formula.xlsx")[0]
    assert "=1+2" in block.text
    assert "未计算" in block.warning
    assert block.location["cell_range"] == "A2:C2"


def test_ragged_csv_reports_coordinate_warning_without_dropping_extra_cell():
    block = parse(b"Model,Value\nCF-100,1,extra\n", "ragged.csv")[0]
    assert "extra" in block.text
    assert block.location["column_end"] == 3
    assert block.warning


def test_pdf_margin_cleaning_keeps_original_spans_for_review():
    pages = [
        Block(
            f"Repeated header\nBody page {i} with actual instructions\nRepeated footer",
            {"type": "pdf", "page": i},
        )
        for i in range(1, 4)
    ]
    cleaned = repeated_page_margins(pages)
    for block, original in zip(cleaned, pages):
        assert block.text == original.text
        parts = chunk([block], overlap=0)
        assert all(
            "Repeated header" not in c["content"]
            and "Repeated footer" not in c["content"]
            for c in parts
        )
        assert any("Repeated header" in c["source_text"] for c in parts)
        for part in parts:
            loc = json.loads(part["location"])
            assert (
                part["source_text"]
                == original.text[loc["block_start"] : loc["block_end"]]
            )
            assert loc["cleaning"]["ignored_spans"]
    assert all(not b.ignored_spans for b in repeated_page_margins(pages[:2]))


def test_real_text_pdf_preserves_page_location_without_ocr():
    from pypdf import PdfWriter
    from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

    writer = PdfWriter()
    page = writer.add_blank_page(width=612, height=792)
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    page[NameObject("/Resources")] = DictionaryObject(
        {NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})}
    )
    stream = DecodedStreamObject()
    stream.set_data(
        b"BT /F1 12 Tf 50 700 Td (CF-100 installation manual: disconnect power before service.) Tj ET"
    )
    page[NameObject("/Contents")] = stream
    output = BytesIO()
    writer.write(output)
    blocks = parse(output.getvalue(), "manual.pdf")
    assert "CF-100" in blocks[0].text
    assert blocks[0].location["page"] == 1
    assert "OCR" not in blocks[0].warning
