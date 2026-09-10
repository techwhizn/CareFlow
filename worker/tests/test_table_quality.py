import json
from io import BytesIO

from careflow.parsing import chunk, parse


def test_markdown_table_repeats_headers_per_row_with_exact_source_offsets():
    source = "# Measurements\nIntro\n| Model | Mass (kg) |\n| --- | ---: |\n| CF-100 | 12 |\n| CF-200 | 14 |\n\nEnd\n"
    blocks = parse(source.encode(), "measurements.md")
    tables = [block for block in blocks if block.location["type"] == "table"]
    assert len(tables) == 2
    assert [block.location["row"] for block in tables] == [2, 3]
    assert tables[1].location["headers"] == ["Model", "Mass (kg)"]
    for block in blocks:
        assert block.text == source[block.location["start"] : block.location["end"]]
    assert all("Mass (kg)" in row["content"] for row in chunk(tables))
    fenced = "```text\n| Example | Value |\n| --- | --- |\n| CF-100 | 12 |\n```\n"
    assert all(
        block.location["type"] == "text" for block in parse(fenced.encode(), "code.md")
    )


def test_csv_labels_units_and_row_identity_survive_native_budget_splits():
    source = "Model,Weight (kg),Notes\nCF-100,12," + "long description " * 60 + "\n"
    blocks = parse(source.encode(), "measurements.csv")
    location = blocks[0].location
    assert location["table_name"] == "measurements.csv"
    assert location["header_annotations"] == [[], ["kg"], []]
    assert location["row_key"] == "2"
    parts = chunk(blocks, target=70, maximum=100, overlap=0)
    assert len(parts) > 1
    for part in parts:
        assert "Weight (kg)" in part["content"]
        loc = json.loads(part["location"])
        assert loc["row"] == 2
        assert loc["block_length"] == len(blocks[0].text)
        assert (
            part["source_text"] == blocks[0].text[loc["block_start"] : loc["block_end"]]
        )


def test_table_quality_marks_empty_duplicate_missing_headers_and_broken_rows():
    blocks = parse(b"Model,,Model\n,,\nCF-100,12,ok,extra\n", "bad.csv")
    assert set(blocks[0].location["quality_codes"]) == {
        "TABLE_HEADER_MISSING",
        "TABLE_HEADER_DUPLICATE",
        "EMPTY_CONTENT",
    }
    assert "TABLE_ROW_BROKEN" in blocks[1].location["quality_codes"]
    assert "extra" in blocks[1].text


def test_word_caption_is_preserved_verbatim_as_table_name():
    from docx import Document

    document = Document()
    document.add_paragraph("Table 2. Power limits", style="Caption")
    table = document.add_table(rows=2, cols=2)
    for cell, value in zip(
        [table.cell(0, 0), table.cell(0, 1), table.cell(1, 0), table.cell(1, 1)],
        ["Model", "Power (W)", "CF-100", "40"],
    ):
        cell.text = value
    output = BytesIO()
    document.save(output)
    block = parse(output.getvalue(), "power.docx")[1]
    assert block.location["table_name"] == "Table 2. Power limits"
    assert block.location["header_annotations"] == [[], ["W"]]
